"""
Регрессионные тесты для подтверждённых критических проблем HydroSphere.

Каждый тест воспроизводит конкретную ошибку из аудита:
    CR-1  data_loader.py:121       UnboundLocalError: 'pd' (универсальный Excel)
    CR-2  frequency.py:114         kritsky_menkel_ppf -> NaN при Cs < 0
    CR-3  trends.py:41-42          mann_kendall_test всегда S=0, p=1.0
    CR-4  flood_hydrograph.py:151  unit_hydrograph не сохраняет объём
    CR-5  homogeneity.py:427       F-test: одностороннее p-value при двусторонней гипотезе
    CR-6  homogeneity.py:418       Welch t-test: pooled df вместо Welch-Satterthwaite

Тесты проверяют математический результат и контракт функций,
не внутренние детали реализации.
"""

import numpy as np
from scipy import stats

# Проектные модули доступны, т.к. pytest добавляет корень проекта в sys.path
from core.hydrorash.flood_hydrograph import hydrograph_convolution, unit_hydrograph
from core.stats.data_loader import parse_hydro_data
from core.stats.frequency import kritsky_menkel_ppf
from core.stats.homogeneity import stationarity_test
from core.stats.trends import mann_kendall_test


# ============================================================
# CR-1: data_loader.py — универсальный Excel (UnboundLocalError)
# ============================================================

def _make_universal_xlsx(path):
    """Универсальный формат: A4='река', строка 5+ содержит река/пост/год/расход."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A4'] = 'река'
    ws['B4'] = 'пост'
    ws['C4'] = 'год'
    ws['D4'] = 'расход'
    rows = [
        ('р.Тест', 'пост1', 2000, 12.5),
        ('р.Тест', 'пост1', 2001, 14.2),
        ('р.Тест', 'пост1', 2002, 11.8),
    ]
    for i, (r, g, y, q) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=r)
        ws.cell(row=i, column=2, value=g)
        ws.cell(row=i, column=3, value=y)
        ws.cell(row=i, column=4, value=q)
    wb.save(path)
    return path


def _make_sheet_xlsx(path):
    """Листовой формат: B1=река, B2=пост, строка 5+: год/расход."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['B1'] = 'р.Лист'
    ws['B2'] = 'пост2'
    rows = [(1990, 10.0), (1991, 12.0), (1992, 11.5), (1993, 13.0)]
    for i, (y, q) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=y)
        ws.cell(row=i, column=2, value=q)
    wb.save(path)
    return path


def test_cr1_universal_format_loads_without_error(tmp_path):
    """CR-1: универсальный формат Excel должен загружаться без UnboundLocalError."""
    path = _make_universal_xlsx(str(tmp_path / 'universal.xlsx'))
    data = parse_hydro_data(path)  # раньше: UnboundLocalError: pd
    assert len(data) == 1
    key = 'р.Тест_пост1'
    assert key in data
    df = data[key]['df']
    assert len(df) == 3
    assert list(df.columns) == ['year', 'Q']
    assert int(df['year'].iloc[0]) == 2000
    assert float(df['Q'].iloc[2]) == 11.8


def test_cr1_sheet_format_still_works(tmp_path):
    """CR-1: листовой формат не должен ломаться после исправления."""
    path = _make_sheet_xlsx(str(tmp_path / 'sheet.xlsx'))
    data = parse_hydro_data(path)
    assert 'р.Лист_пост2' in data
    df = data['р.Лист_пост2']['df']
    assert len(df) == 4
    assert int(df['year'].min()) == 1990


# ============================================================
# CR-2: kritsky_menkel_ppf — NaN при Cs < 0
# ============================================================

def test_cr2_kritsky_menkel_negative_cs_no_nan():
    """CR-2: для Cs < 0 квантили К-М не должны содержать NaN (обрезание до 0 допустимо)."""
    probs = np.array([0.001, 0.01, 0.05, 0.1, 0.5, 0.9, 0.99])
    for cv in (0.1, 0.3, 0.5):
        for cs in (-0.2, -0.5, -1.0):
            q = kritsky_menkel_ppf(probs, mean=100.0, cv=cv, cs=cs)
            assert q.shape == probs.shape
            assert not np.any(np.isnan(q)), f'NaN при cv={cv}, cs={cs}: {q}'
            assert np.all(q >= 0), f'Отрицательные квантили (без обрезания) при cv={cv}, cs={cs}: {q}'
            assert np.all(np.isfinite(q)), f'Не-конечные значения при cv={cv}, cs={cs}: {q}'


def test_cr2_kritsky_menkel_negative_cs_ordering():
    """CR-2: квантили при Cs<0 должны быть монотонно невозрастающими по P (обеспеченность)."""
    probs = np.array([0.001, 0.01, 0.05, 0.1, 0.3, 0.5, 0.7, 0.9, 0.99])
    q = kritsky_menkel_ppf(probs, mean=50.0, cv=0.4, cs=-0.6)
    # Большая обеспеченность P -> меньший квантиль X
    diffs = np.diff(q)
    assert np.all(diffs <= 1e-9 + 1e-6 * np.abs(q[:-1])), f'Немонотонно: {q}'


def test_cr2_kritsky_menkel_median_near_mean_for_cs_negative():
    """CR-2: медиана (P=0.5) при Cs<0 близка к среднему при малом Cv."""
    q = kritsky_menkel_ppf(np.array([0.5]), mean=100.0, cv=0.1, cs=-0.3)
    # Для слабой асимметрии медиана отклоняется от среднего незначительно (Cv=0.1)
    assert abs(float(q[0]) - 100.0) < 8.0, f'Медиана далеко от среднего: {q[0]}'


def test_cr2_kritsky_menkel_positive_cs_preserved():
    """CR-2: поведение для Cs>0 не должно измениться (сохранение результатов)."""
    probs = np.array([0.01, 0.05, 0.5, 0.95])
    # Эталон: текущая (рабочая) ветка Cs>0 — гамма с q=1-p
    for cv, cs in ((0.2, 0.5), (0.3, 1.0)):
        mean = 100.0
        alpha = 4.0 / cs ** 2
        beta = mean * cv * cs / 2.0
        a0 = mean * (1.0 - 2.0 * cv / cs)
        expected = a0 + stats.gamma.ppf(1 - probs, a=alpha, scale=beta)
        got = kritsky_menkel_ppf(probs, mean=mean, cv=cv, cs=cs)
        np.testing.assert_allclose(got, expected, rtol=1e-9)


# ============================================================
# CR-3: mann_kendall_test — S=0, p=1.0
# ============================================================

def _mk_expected_s(data):
    """Независимый эталон S = sum_{i<j} sign(x_j - x_i)."""
    n = len(data)
    s = 0
    for i in range(n):
        for j in range(i + 1, n):
            s += np.sign(data[j] - data[i])
    return s


def _mk_expected_var(n, data):
    """Дисперсия S с поправкой на связи (та же формула, что в функции)."""
    _, counts = np.unique(data, return_counts=True)
    ties = np.sum(counts * (counts - 1) * (2 * counts + 5))
    var = (n * (n - 1) * (2 * n + 5) - ties) / 18.0
    return max(var, 1e-10)


def test_cr3_increasing_series():
    """CR-3: [1,2,3,4,5] -> S=+10, положительный тренд, p < 0.05."""
    data = np.array([1.0, 2.0, 3.0, 4.0, 5.0])
    res = mann_kendall_test(data)
    assert res['statistic'] == 10.0
    n = len(data)
    var = _mk_expected_var(n, data)
    z_exp = (10.0 - 1) / np.sqrt(var)
    p_exp = 2 * (1 - stats.norm.cdf(abs(z_exp)))
    assert abs(res['z'] - z_exp) < 1e-9
    assert abs(res['p_value'] - p_exp) < 1e-9
    assert res['trend'] == 'Рост'
    assert res['significant'] is True


def test_cr3_decreasing_series():
    """CR-3: [5,4,3,2,1] -> S=-10, отрицательный тренд."""
    data = np.array([5.0, 4.0, 3.0, 2.0, 1.0])
    res = mann_kendall_test(data)
    assert res['statistic'] == -10.0
    n = len(data)
    var = _mk_expected_var(n, data)
    z_exp = (-10.0 + 1) / np.sqrt(var)
    p_exp = 2 * (1 - stats.norm.cdf(abs(z_exp)))
    assert abs(res['z'] - z_exp) < 1e-9
    assert abs(res['p_value'] - p_exp) < 1e-9
    assert res['trend'] == 'Снижение'


def test_cr3_constant_series():
    """CR-3: [1,1,1,1,1] -> S=0, тренд отсутствует."""
    data = np.array([1.0, 1.0, 1.0, 1.0, 1.0])
    res = mann_kendall_test(data)
    assert res['statistic'] == 0.0
    assert res['z'] == 0.0
    assert res['p_value'] == 1.0
    assert res['trend'] == 'Тренд отсутствует'
    assert res['significant'] is False


def test_cr3_mixed_series():
    """CR-3: [1,3,2,5,4] -> S=+6 (сверка с ручным расчётом)."""
    data = np.array([1.0, 3.0, 2.0, 5.0, 4.0])
    res = mann_kendall_test(data)
    s_manual = _mk_expected_s(data)
    assert res['statistic'] == s_manual == 6.0
    n = len(data)
    var = _mk_expected_var(n, data)
    z_exp = (6.0 - 1) / np.sqrt(var)
    p_exp = 2 * (1 - stats.norm.cdf(abs(z_exp)))
    assert abs(res['z'] - z_exp) < 1e-9
    assert abs(res['p_value'] - p_exp) < 1e-9
    assert res['trend'] == 'Рост'


def test_cr3_matches_independent_loop_for_random_series():
    """CR-3: векторная формула S совпадает с независимым циклом на случайном ряду."""
    rng = np.random.default_rng(42)
    data = rng.normal(10, 2, 40)
    res = mann_kendall_test(data)
    assert res['statistic'] == float(_mk_expected_s(data))


# ============================================================
# CR-4: unit_hydrograph — сохранение объёма
# ============================================================

def test_cr4_unit_hydrograph_volume_conservation():
    """CR-4: 100 км2 x 1 мм -> объём = 100 000 м3 (в пределах дискретизации)."""
    F = 100.0
    T_peak, T_base, shape = 5.0, 24.0, 3.5
    dt = 0.5  # мелкий шаг для малой погрешности численного интегрирования
    uh = unit_hydrograph(T_peak=T_peak, T_base=T_base, F_km2=F, shape=shape, dt=dt)
    expected = 1000.0 * F  # 1 мм слоя стока
    ratio = uh['volume_m3'] / expected
    assert 0.97 <= ratio <= 1.03, \
        f'Объём единичного гидрографа {uh["volume_m3"]:.0f} м3, ожидалось {expected:.0f} (ratio={ratio:.3f})'


def test_cr4_unit_hydrograph_volume_linear_in_area():
    """CR-4: удвоение площади -> удвоение объёма (линейность)."""
    args = dict(T_peak=5.0, T_base=24.0, shape=3.5, dt=0.5)
    uh1 = unit_hydrograph(F_km2=100.0, **args)
    uh2 = unit_hydrograph(F_km2=200.0, **args)
    ratio = uh2['volume_m3'] / uh1['volume_m3']
    assert abs(ratio - 2.0) < 0.02, f'Линейность по площади нарушена: ratio={ratio:.3f}'


def test_cr4_convolution_scales_with_rainfall():
    """CR-4: свёртка с двойным слоем осадков даёт двойной объём."""
    args = dict(T_peak=5.0, T_base=24.0, F_km2=100.0, shape=3.5, dt=0.5)
    uh = unit_hydrograph(**args)
    q1 = hydrograph_convolution(np.array(uh['Q_m3_s']), np.array([1.0]), dt=0.5)
    q2 = hydrograph_convolution(np.array(uh['Q_m3_s']), np.array([2.0]), dt=0.5)
    v1 = np.trapezoid(q1, dx=0.5) * 3600
    v2 = np.trapezoid(q2, dx=0.5) * 3600
    assert abs(v2 / v1 - 2.0) < 1e-6


# ============================================================
# CR-5: F-test — двустороннее p-value
# ============================================================

def test_cr5_ftest_two_sided_pvalue():
    """CR-5: p-value F-теста должно быть двусторонним (равенство дисперсий)."""
    # Контролируемый пример: разные дисперсии, разный размер выборок
    part1 = np.arange(1, 11, dtype=float) * 2.0        # n1=10
    part2 = np.arange(1, 26, dtype=float) * 5.0        # n2=25, дисперсия >> дисперсии part1
    data = np.concatenate([part1, part2])
    years = np.arange(len(data))
    # Разделение пополам -> части получаются неравными по дисперсии
    result = stationarity_test(data, years=years, alpha=0.05)

    # Независимый расчёт по scipy.sf: двусторонний p = 2 * sf(F)
    var1 = np.var(data[:len(data) // 2], ddof=1)
    var2 = np.var(data[len(data) // 2:], ddof=1)
    f_stat = max(var1, var2) / min(var1, var2)
    n1, n2 = len(data) // 2, len(data) - len(data) // 2
    df_num = (n1 - 1) if var1 >= var2 else (n2 - 1)
    df_den = (n2 - 1) if var1 >= var2 else (n1 - 1)
    p_two_sided = 2.0 * stats.f.sf(f_stat, df_num, df_den)

    assert abs(result['f_test']['f_stat'] - round(f_stat, 4)) < 1e-6
    assert abs(result['f_test']['p_value'] - round(p_two_sided, 4)) < 1e-6
    # p_value не должен быть равен одностороннему (это была ошибка)
    p_one_sided = stats.f.sf(f_stat, df_num, df_den)
    assert abs(result['f_test']['p_value'] - round(p_one_sided, 4)) > 1e-6


def test_cr5_ftest_significance_consistent_with_pvalue():
    """CR-5: significant = (p_value < alpha) должно согласовываться с p-value."""
    rng = np.random.default_rng(7)
    # явно разные дисперсии — тест должен быть значимым
    part1 = rng.normal(10, 1.0, 40)
    part2 = rng.normal(10, 3.0, 40)
    data = np.concatenate([part1, part2])
    years = np.arange(len(data))
    result = stationarity_test(data, years=years, alpha=0.05)
    p = result['f_test']['p_value']
    assert result['f_test']['significant'] == (p < 0.05)


# ============================================================
# CR-6: Welch t-test — Welch-Satterthwaite df
# ============================================================

def test_cr6_welch_df_consistent_with_scipy():
    """CR-6: критическое значение t должно использовать Welch-Satterthwaite df."""
    rng = np.random.default_rng(123)
    n1, n2 = 8, 24  # разные размеры выборок
    part1 = rng.normal(100, 15.0, n1)   # большая дисперсия
    part2 = rng.normal(103, 4.0, n2)    # малая дисперсия
    data = np.concatenate([part1, part2])
    years = np.arange(len(data))
    # Явное разделение по году: часть1 = первые 8, часть2 = остальные 24
    result = stationarity_test(data, years=years, split_year=8, alpha=0.05)
    assert result['split_info']['n1'] == 8
    assert result['split_info']['n2'] == 24

    # Welch-Satterthwaite df по формуле (независимо)
    s1 = np.var(part1, ddof=1)
    s2 = np.var(part2, ddof=1)
    welch_df = (s1 / n1 + s2 / n2) ** 2 / (
        (s1 / n1) ** 2 / (n1 - 1) + (s2 / n2) ** 2 / (n2 - 1)
    )
    t_crit_expected = stats.t.ppf(1 - 0.05 / 2, welch_df)

    assert abs(result['t_test']['t_critical'] - round(t_crit_expected, 4)) < 1e-6, \
        f't_critical={result["t_test"]["t_critical"]}, ожидалось {round(t_crit_expected, 4)} (df={welch_df:.2f})'


def test_cr6_welch_df_not_pooled():
    """CR-6: df не должен быть равен n1+n2-2 при разных дисперсиях/размерах."""
    rng = np.random.default_rng(123)
    n1, n2 = 8, 24
    part1 = rng.normal(100, 15.0, n1)
    part2 = rng.normal(103, 4.0, n2)
    data = np.concatenate([part1, part2])
    years = np.arange(len(data))
    result = stationarity_test(data, years=years, split_year=8, alpha=0.05)
    pooled_t_crit = stats.t.ppf(1 - 0.05 / 2, n1 + n2 - 2)
    assert abs(result['t_test']['t_critical'] - round(pooled_t_crit, 4)) > 1e-6


def test_cr6_welch_equal_variances_matches_pooled():
    """CR-6: при равных дисперсиях Welch df ~ pooled df (совместимость)."""
    rng = np.random.default_rng(5)
    n1, n2 = 20, 20
    part1 = rng.normal(100, 5.0, n1)
    part2 = rng.normal(102, 5.0, n2)
    data = np.concatenate([part1, part2])
    years = np.arange(len(data))
    result = stationarity_test(data, years=years, alpha=0.05)
    s1 = np.var(part1, ddof=1)
    s2 = np.var(part2, ddof=1)
    welch_df = (s1 / n1 + s2 / n2) ** 2 / (
        (s1 / n1) ** 2 / (n1 - 1) + (s2 / n2) ** 2 / (n2 - 1)
    )
    t_crit_expected = stats.t.ppf(1 - 0.05 / 2, welch_df)
    assert abs(result['t_test']['t_critical'] - round(t_crit_expected, 4)) < 1e-6
